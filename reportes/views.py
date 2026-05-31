from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Count, Q, Sum, Avg, F
from django.utils import timezone
from datetime import datetime, timedelta
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.core.files.base import ContentFile
from django.contrib.admin.models import LogEntry, ADDITION, CHANGE, DELETION
import json
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from .models import Reporte, ConfiguracionReporte
from personal.models import RegistroAcceso, Persona
from dispositivos.models import Dispositivo
from sedes.models import Sede, Area

# ========== IMPORTACIONES SEGURAS ==========
try:
    from tickets.models import Ticket
    TICKETS_AVAILABLE = True
except ImportError:
    Ticket = None
    TICKETS_AVAILABLE = False

try:
    from alertas.models import Alerta
    ALERTAS_AVAILABLE = True
except ImportError:
    Alerta = None
    ALERTAS_AVAILABLE = False


@login_required
def index(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    dispositivo_id = request.GET.get('dispositivo_id')
    tipo_acceso = request.GET.get('tipo_acceso')
    
    registros = RegistroAcceso.objects.select_related('persona', 'dispositivo').all()
    
    if fecha_inicio:
        try:
            fecha_inicio_dt = datetime.strptime(fecha_inicio, '%Y-%m-%d')
            registros = registros.filter(fecha_hora__date__gte=fecha_inicio_dt)
        except ValueError:
            pass
    
    if fecha_fin:
        try:
            fecha_fin_dt = datetime.strptime(fecha_fin, '%Y-%m-%d')
            registros = registros.filter(fecha_hora__date__lte=fecha_fin_dt)
        except ValueError:
            pass
    
    if dispositivo_id and dispositivo_id != '':
        registros = registros.filter(dispositivo_id=dispositivo_id)
    
    if tipo_acceso and tipo_acceso != '':
        registros = registros.filter(tipo_acceso=tipo_acceso)
    
    total_registros = registros.count()
    accesos_exitosos = registros.filter(tipo_acceso='exitoso').count()
    accesos_fallidos = registros.filter(tipo_acceso='fallido').count()
    
    hace_7_dias = timezone.now() - timedelta(days=7)
    registros_semana = registros.filter(fecha_hora__gte=hace_7_dias).count()
    
    dispositivos = Dispositivo.objects.all()
    registros = registros.order_by('-fecha_hora')
    
    usuarios_activos = Persona.objects.filter(activo=True).count()
    sedes_total = Sede.objects.count()
    areas_total = Area.objects.count()
    dispositivos_total = Dispositivo.objects.count()
    
    tickets_pendientes = 0
    if TICKETS_AVAILABLE and Ticket:
        if hasattr(Ticket, 'estado'):
            tickets_pendientes = Ticket.objects.filter(estado='PENDIENTE').count()
    
    alertas_criticas = 0
    if ALERTAS_AVAILABLE and Alerta:
        alertas_criticas = Alerta.objects.filter(nivel='CRITICO', resuelta=False).count()
    
    context = {
        'registros': registros,
        'dispositivos': dispositivos,
        'total_registros': total_registros,
        'accesos_exitosos': accesos_exitosos,
        'accesos_fallidos': accesos_fallidos,
        'registros_semana': registros_semana,
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'dispositivo_id': dispositivo_id or '',
        'tipo_acceso': tipo_acceso or '',
        'usuarios_activos': usuarios_activos,
        'sedes_total': sedes_total,
        'areas_total': areas_total,
        'dispositivos_total': dispositivos_total,
        'tickets_pendientes': tickets_pendientes,
        'alertas_criticas': alertas_criticas,
    }
    
    return render(request, 'reportes/index.html', context)


@login_required
def reportes_guardados(request):
    reportes = Reporte.objects.all()
    return render(request, 'reportes/guardados.html', {'reportes': reportes})


@login_required
def detalle_reporte(request, reporte_id):
    reporte = get_object_or_404(Reporte, id=reporte_id)
    return render(request, 'reportes/detalle.html', {'reporte': reporte})


@csrf_exempt
@login_required
def generar_reporte(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)
    
    try:
        data = json.loads(request.body)
        
        titulo = data.get('titulo', '')
        tipo = data.get('tipo', 'GENERAL')
        fecha_inicio = data.get('fecha_inicio')
        fecha_fin = data.get('fecha_fin')
        
        if not fecha_inicio or not fecha_fin:
            return JsonResponse({'success': False, 'message': 'Rango de fechas requerido'}, status=400)
        
        reporte = Reporte.objects.create(
            titulo=titulo or f"Reporte {tipo} - {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            tipo=tipo,
            estado='PROCESANDO',
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
        )
        
        procesar_reporte_completo(reporte.id)
        
        return JsonResponse({
            'success': True,
            'message': 'Reporte generado correctamente',
            'reporte_id': reporte.id
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


def procesar_reporte_completo(reporte_id):
    reporte = get_object_or_404(Reporte, id=reporte_id)
    
    try:
        if isinstance(reporte.fecha_inicio, str):
            fecha_inicio = datetime.strptime(reporte.fecha_inicio, '%Y-%m-%d')
            fecha_fin = datetime.strptime(reporte.fecha_fin, '%Y-%m-%d')
        else:
            fecha_inicio = reporte.fecha_inicio
            fecha_fin = reporte.fecha_fin
        
        registros = RegistroAcceso.objects.filter(
            fecha_hora__date__gte=fecha_inicio,
            fecha_hora__date__lte=fecha_fin
        )
        
        reporte.total_registros = registros.count()
        reporte.accesos_exitosos = registros.filter(tipo_acceso='exitoso').count()
        reporte.accesos_fallidos = registros.filter(tipo_acceso='fallido').count()
        
        reporte.usuarios_nuevos = Persona.objects.filter(
            fecha_creacion__date__gte=fecha_inicio,
            fecha_creacion__date__lte=fecha_fin
        ).count()
        
        reporte.usuarios_activos = Persona.objects.filter(activo=True).count()
        reporte.usuarios_inactivos = Persona.objects.filter(activo=False).count()
        
        reporte.sedes_nuevas = Sede.objects.filter(
            fecha_creacion__date__gte=fecha_inicio,
            fecha_creacion__date__lte=fecha_fin
        ).count()
        
        if hasattr(Dispositivo, 'estado'):
            reporte.dispositivos_activos = Dispositivo.objects.filter(estado='activo').count()
            reporte.dispositivos_inactivos = Dispositivo.objects.filter(estado='inactivo').count()
        else:
            reporte.dispositivos_activos = Dispositivo.objects.filter(activo=True).count()
            reporte.dispositivos_inactivos = Dispositivo.objects.filter(activo=False).count()
        
        if ALERTAS_AVAILABLE and Alerta:
            reporte.alertas_totales = Alerta.objects.filter(
                fecha_hora__date__gte=fecha_inicio,
                fecha_hora__date__lte=fecha_fin
            ).count()
            reporte.alertas_criticas = Alerta.objects.filter(nivel='CRITICO', resuelta=False).count()
        
        if reporte.total_registros > 0:
            reporte.tasa_exito_global = (reporte.accesos_exitosos / reporte.total_registros) * 100
        
        reporte.datos_json = {
            'accesos': {
                'total': reporte.total_registros,
                'exitosos': reporte.accesos_exitosos,
                'fallidos': reporte.accesos_fallidos,
            },
            'usuarios': {
                'nuevos': reporte.usuarios_nuevos,
                'activos': reporte.usuarios_activos,
                'inactivos': reporte.usuarios_inactivos,
            },
            'alertas': {
                'totales': reporte.alertas_totales,
                'criticas': reporte.alertas_criticas,
            },
            'rendimiento': {
                'tasa_exito': round(reporte.tasa_exito_global, 2),
            }
        }
        
        reporte.fecha_generacion = timezone.now()
        reporte.estado = 'COMPLETADO'
        reporte.save()
        
        generar_archivo_reporte_completo(reporte)
        
    except Exception as e:
        reporte.estado = 'ERROR'
        reporte.save()
        raise e


def generar_archivo_reporte_completo(reporte):
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(['=== INFORME DEL SISTEMA BIOMETRIKA ==='])
    writer.writerow(['Título', reporte.titulo])
    writer.writerow(['Tipo', reporte.get_tipo_display()])
    writer.writerow(['Período', f"{reporte.fecha_inicio} al {reporte.fecha_fin}"])
    writer.writerow(['Fecha Generación', reporte.fecha_generacion.strftime('%Y-%m-%d %H:%M:%S') if reporte.fecha_generacion else 'N/A'])
    writer.writerow([])
    
    writer.writerow(['=== 1. ACCESOS BIOMÉTRICOS ==='])
    writer.writerow(['Total Registros', reporte.total_registros])
    writer.writerow(['Accesos Exitosos', reporte.accesos_exitosos])
    writer.writerow(['Accesos Fallidos', reporte.accesos_fallidos])
    writer.writerow(['Tasa de Éxito', f"{reporte.tasa_exito_global:.2f}%"])
    writer.writerow([])
    
    writer.writerow(['=== 2. GESTIÓN DE USUARIOS ==='])
    writer.writerow(['Usuarios Nuevos', reporte.usuarios_nuevos])
    writer.writerow(['Usuarios Activos', reporte.usuarios_activos])
    writer.writerow(['Usuarios Inactivos', reporte.usuarios_inactivos])
    writer.writerow([])
    
    writer.writerow(['=== 3. ALERTAS ==='])
    writer.writerow(['Alertas Totales', reporte.alertas_totales])
    writer.writerow(['Alertas Críticas', reporte.alertas_criticas])
    writer.writerow([])
    
    writer.writerow(['=== 4. RESUMEN EJECUTIVO ==='])
    salud = 'Buena' if reporte.tasa_exito_global > 90 else 'Regular' if reporte.tasa_exito_global > 70 else 'Crítica'
    writer.writerow(['Salud General del Sistema', salud])
    writer.writerow(['Total de Usuarios en el Sistema', reporte.usuarios_activos])
    
    csv_content = output.getvalue()
    reporte.archivo.save(
        f'reporte_completo_{reporte.id}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv',
        ContentFile(csv_content.encode('utf-8'))
    )
    output.close()


@csrf_exempt
@login_required
def eliminar_reporte(request, reporte_id):
    if request.method != 'DELETE':
        return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)
    
    try:
        reporte = get_object_or_404(Reporte, id=reporte_id)
        if reporte.archivo:
            reporte.archivo.delete()
        reporte.delete()
        return JsonResponse({'success': True, 'message': 'Reporte eliminado correctamente'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
def exportar_reporte(request, reporte_id, formato):
    reporte = get_object_or_404(Reporte, id=reporte_id)
    
    if not reporte.archivo:
        return JsonResponse({'success': False, 'message': 'El reporte no tiene archivo asociado'}, status=400)
    
    if formato == 'csv':
        response = HttpResponse(reporte.archivo.read(), content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{reporte.archivo.name}"'
        return response
    
    elif formato == 'json':
        response = HttpResponse(json.dumps(reporte.datos_json, ensure_ascii=False, indent=2), content_type='application/json')
        response['Content-Disposition'] = f'attachment; filename="reporte_{reporte.id}.json"'
        return response
    
    elif formato in ['excel', 'xlsx']:
        csv_content = reporte.archivo.read().decode('utf-8')
        reader = csv.reader(io.StringIO(csv_content))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte General"
        
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 1 or (value and str(value).startswith('===')):
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                    cell.font = header_font
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reporte_{reporte.id}.xlsx"'
        wb.save(response)
        return response
    
    elif formato == 'txt':
        content = reporte.archivo.read().decode('utf-8')
        response = HttpResponse(content, content_type='text/plain')
        response['Content-Disposition'] = f'attachment; filename="reporte_{reporte.id}.txt"'
        return response
    
    return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)


@login_required
def exportar_datos(request, formato):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    dispositivo_id = request.GET.get('dispositivo_id')
    tipo_acceso = request.GET.get('tipo_acceso')
    
    registros = RegistroAcceso.objects.select_related('persona', 'dispositivo').all()
    
    if fecha_inicio:
        try:
            fecha_inicio_dt = datetime.strptime(fecha_inicio, '%Y-%m-%d')
            registros = registros.filter(fecha_hora__date__gte=fecha_inicio_dt)
        except:
            pass
    
    if fecha_fin:
        try:
            fecha_fin_dt = datetime.strptime(fecha_fin, '%Y-%m-%d')
            registros = registros.filter(fecha_hora__date__lte=fecha_fin_dt)
        except:
            pass
    
    if dispositivo_id and dispositivo_id != '':
        registros = registros.filter(dispositivo_id=dispositivo_id)
    
    if tipo_acceso and tipo_acceso != '':
        registros = registros.filter(tipo_acceso=tipo_acceso)
    
    if formato == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="reporte_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv"'
        response.write('\ufeff')
        
        writer = csv.writer(response)
        writer.writerow(['ID', 'Usuario', 'Documento', 'Dispositivo', 'Fecha/Hora', 'Tipo Acceso'])
        
        for r in registros:
            writer.writerow([
                r.id,
                r.persona.nombres + ' ' + r.persona.apellidos if r.persona else 'N/A',
                r.persona.numero_documento if r.persona else 'N/A',
                r.dispositivo.nombre if r.dispositivo else 'N/A',
                r.fecha_hora.strftime('%Y-%m-%d %H:%M:%S') if r.fecha_hora else '',
                r.tipo_acceso
            ])
        return response
    
    elif formato == 'json':
        data = []
        for r in registros:
            data.append({
                'id': r.id,
                'usuario': r.persona.nombres + ' ' + r.persona.apellidos if r.persona else None,
                'documento': r.persona.numero_documento if r.persona else None,
                'dispositivo': r.dispositivo.nombre if r.dispositivo else None,
                'fecha_hora': r.fecha_hora.strftime('%Y-%m-%d %H:%M:%S') if r.fecha_hora else None,
                'tipo_acceso': r.tipo_acceso
            })
        response = HttpResponse(json.dumps(data, ensure_ascii=False, indent=2), content_type='application/json')
        response['Content-Disposition'] = f'attachment; filename="reporte_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json"'
        return response
    
    elif formato == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = "Registros"
        
        headers = ['ID', 'Usuario', 'Documento', 'Dispositivo', 'Fecha/Hora', 'Tipo Acceso']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        
        for row, r in enumerate(registros, 2):
            ws.cell(row=row, column=1, value=r.id)
            ws.cell(row=row, column=2, value=r.persona.nombres + ' ' + r.persona.apellidos if r.persona else 'N/A')
            ws.cell(row=row, column=3, value=r.persona.numero_documento if r.persona else 'N/A')
            ws.cell(row=row, column=4, value=r.dispositivo.nombre if r.dispositivo else 'N/A')
            ws.cell(row=row, column=5, value=r.fecha_hora.strftime('%Y-%m-%d %H:%M:%S') if r.fecha_hora else '')
            ws.cell(row=row, column=6, value=r.tipo_acceso)
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reporte_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        wb.save(response)
        return response
    
    return JsonResponse({'error': 'Formato no soportado'}, status=400)


@login_required
def estadisticas_api(request):
    dias = int(request.GET.get('dias', 30))
    fecha_inicio = timezone.now() - timedelta(days=dias)
    
    datos_por_dia = {}
    registros = RegistroAcceso.objects.filter(fecha_hora__gte=fecha_inicio)
    
    for registro in registros:
        dia_key = registro.fecha_hora.strftime('%Y-%m-%d')
        if dia_key not in datos_por_dia:
            datos_por_dia[dia_key] = {'total': 0, 'exitosos': 0, 'fallidos': 0}
        
        datos_por_dia[dia_key]['total'] += 1
        if registro.tipo_acceso == 'exitoso':
            datos_por_dia[dia_key]['exitosos'] += 1
        else:
            datos_por_dia[dia_key]['fallidos'] += 1
    
    datos_dia = []
    for dia, stats in sorted(datos_por_dia.items()):
        datos_dia.append({
            'dia': dia,
            'total': stats['total'],
            'exitosos': stats['exitosos'],
            'fallidos': stats['fallidos']
        })
    
    tickets_data = {'pendientes': 0}
    if TICKETS_AVAILABLE and Ticket and hasattr(Ticket, 'estado'):
        tickets_data['pendientes'] = Ticket.objects.filter(estado='PENDIENTE').count()
    
    if hasattr(Dispositivo, 'estado'):
        dispositivos_activos = Dispositivo.objects.filter(estado='activo').count()
        dispositivos_inactivos = Dispositivo.objects.filter(estado='inactivo').count()
    else:
        dispositivos_activos = Dispositivo.objects.filter(activo=True).count()
        dispositivos_inactivos = Dispositivo.objects.filter(activo=False).count()
    
    dispositivos_data = {
        'activos': dispositivos_activos,
        'inactivos': dispositivos_inactivos,
    }
    
    alertas_data = {'criticas': 0, 'problematicas': 0, 'estables': 0}
    if ALERTAS_AVAILABLE and Alerta:
        alertas_data['criticas'] = Alerta.objects.filter(nivel='CRITICO', resuelta=False).count()
        alertas_data['problematicas'] = Alerta.objects.filter(nivel='PROBLEMA', resuelta=False).count()
        alertas_data['estables'] = Alerta.objects.filter(nivel='ESTABLE', resuelta=False).count()
    
    total_general = registros.count()
    exitosos_general = registros.filter(tipo_acceso='exitoso').count()
    tasa_exito = (exitosos_general / total_general * 100) if total_general > 0 else 0
    
    return JsonResponse({
        'success': True,
        'datos_dia': datos_dia,
        'tasa_exito': round(tasa_exito, 2),
        'total_registros': total_general,
        'tickets': tickets_data,
        'dispositivos': dispositivos_data,
        'alertas': alertas_data,
        'usuarios_activos': Persona.objects.filter(activo=True).count(),
    })