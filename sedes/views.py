from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
import json
import csv
import io
from openpyxl import Workbook, load_workbook
import openpyxl.styles
from .models import Sede, Area

@login_required
def lista_sedes(request):
    sedes = Sede.objects.all()
    return render(request, 'sedes/lista.html', {'sedes': sedes})

# ==================== API SEDES ====================

@csrf_exempt
def api_sedes(request):
    if request.method == 'GET':
        try:
            sedes = Sede.objects.all()
            data = []
            for sede in sedes:
                data.append({
                    'id': sede.id,
                    'nombre': sede.nombre,
                    'codigo_unico': sede.codigo_unico,
                    'direccion': sede.direccion,
                    'activo': sede.activo,
                    'tipo': 'sede',
                    'dispositivo_biometrico': sede.dispositivo_biometrico,
                    'nivel_seguridad': sede.nivel_seguridad
                })
            return JsonResponse(data, safe=False)
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    elif request.method == 'POST':
        try:
            data = json.loads(request.body)
            nombre = data.get('nombre', '').strip()
            direccion = data.get('direccion', '').strip()
            dispositivo = data.get('dispositivo_biometrico', 'huella')
            nivel_seguridad = data.get('nivel_seguridad', 'bajo')

            if not nombre:
                return JsonResponse({'success': False, 'message': 'Nombre de sede requerido'}, status=400)

            count = Sede.objects.count() + 1
            codigo_unico = f"SED-{count:03d}"

            sede = Sede.objects.create(
                nombre=nombre,
                direccion=direccion,
                codigo_unico=codigo_unico,
                activo=True,
                dispositivo_biometrico=dispositivo,
                nivel_seguridad=nivel_seguridad
            )

            return JsonResponse({
                'success': True,
                'message': 'Sede creada correctamente',
                'id': sede.id,
                'sede': {
                    'id': sede.id,
                    'nombre': sede.nombre,
                    'codigo_unico': sede.codigo_unico,
                    'direccion': sede.direccion,
                    'dispositivo_biometrico': sede.dispositivo_biometrico,
                    'nivel_seguridad': sede.nivel_seguridad
                }
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)


@csrf_exempt
def api_sede_detail(request, sede_id):
    try:
        sede = Sede.objects.get(id=sede_id)
    except Sede.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Sede no encontrada'}, status=404)

    if request.method == 'GET':
        return JsonResponse({
            'id': sede.id,
            'nombre': sede.nombre,
            'codigo_unico': sede.codigo_unico,
            'direccion': sede.direccion,
            'activo': sede.activo,
            'tipo': 'sede',
            'dispositivo_biometrico': sede.dispositivo_biometrico,
            'nivel_seguridad': sede.nivel_seguridad
        })

    elif request.method == 'PUT':
        try:
            data = json.loads(request.body)
            sede.nombre = data.get('nombre', sede.nombre).strip()
            sede.direccion = data.get('direccion', sede.direccion).strip()
            sede.dispositivo_biometrico = data.get('dispositivo_biometrico', sede.dispositivo_biometrico)
            sede.nivel_seguridad = data.get('nivel_seguridad', sede.nivel_seguridad)
            sede.save()
            return JsonResponse({'success': True, 'message': 'Sede actualizada correctamente'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    elif request.method == 'DELETE':
        try:
            areas_eliminadas = Area.objects.filter(sede=sede).count()
            Area.objects.filter(sede=sede).delete()
            sede.delete()
            return JsonResponse({
                'success': True,
                'message': f'Sede eliminada. Se eliminaron {areas_eliminadas} áreas asociadas.'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)


# ==================== API AREAS ====================

@csrf_exempt
def api_areas(request):
    if request.method == 'GET':
        try:
            areas = Area.objects.select_related('sede').all()
            data = []
            for area in areas:
                data.append({
                    'id': area.id,
                    'sede_id': area.sede.id,
                    'sede_nombre': area.sede.nombre,
                    'nombre': area.nombre,
                    'piso': area.piso,
                    'codigo_acceso': area.codigo_acceso,
                    'dispositivo_biometrico': area.dispositivo_biometrico,
                    'nivel_seguridad': area.nivel_seguridad,
                    'tipo': 'area'
                })
            return JsonResponse(data, safe=False)
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    elif request.method == 'POST':
        try:
            data = json.loads(request.body)
            sede_id = data.get('sede_id')
            nombre = data.get('nombre', '').strip()
            piso = data.get('piso', 1)
            codigo_acceso = data.get('codigo_acceso', '').strip()
            dispositivo = data.get('dispositivo_biometrico', 'huella')
            nivel_seguridad = data.get('nivel_seguridad', 'bajo')

            if not sede_id or not nombre:
                return JsonResponse({'success': False, 'message': 'Faltan campos requeridos'}, status=400)

            try:
                sede = Sede.objects.get(id=sede_id)
            except Sede.DoesNotExist:
                return JsonResponse({'success': False, 'message': 'Sede no encontrada'}, status=404)

            # Validar nivel_seguridad
            nivel_seguridad = str(nivel_seguridad).strip().lower()
            if nivel_seguridad not in ['bajo', 'medio', 'alto']:
                nivel_seguridad = 'bajo'

            # Validar dispositivo
            if dispositivo not in ['huella', 'Tarjeta', 'PIN']:
                dispositivo = 'huella'

            count = Area.objects.filter(sede=sede).count() + 1
            if not codigo_acceso:
                codigo_acceso = f"ARE-{sede.id:03d}-{count:03d}"

            area = Area.objects.create(
                sede=sede,
                nombre=nombre,
                piso=int(piso),
                codigo_acceso=codigo_acceso,
                dispositivo_biometrico=dispositivo,
                nivel_seguridad=nivel_seguridad
            )

            return JsonResponse({
                'success': True,
                'message': 'Área creada correctamente',
                'id': area.id,
                'area': {
                    'id': area.id,
                    'nombre': area.nombre,
                    'sede_id': area.sede.id,
                    'sede_nombre': area.sede.nombre,
                    'piso': area.piso,
                    'codigo_acceso': area.codigo_acceso,
                    'dispositivo_biometrico': area.dispositivo_biometrico,
                    'nivel_seguridad': area.nivel_seguridad,
                    'tipo': 'area'
                }
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)


@csrf_exempt
def api_area_detail(request, area_id):
    try:
        area = Area.objects.select_related('sede').get(id=area_id)
    except Area.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Área no encontrada'}, status=404)

    if request.method == 'GET':
        return JsonResponse({
            'id': area.id,
            'sede_id': area.sede.id,
            'sede_nombre': area.sede.nombre,
            'nombre': area.nombre,
            'piso': area.piso,
            'codigo_acceso': area.codigo_acceso,
            'dispositivo_biometrico': area.dispositivo_biometrico,
            'nivel_seguridad': area.nivel_seguridad,
            'tipo': 'area'
        })

    elif request.method == 'PUT':
        try:
            data = json.loads(request.body)
            area.nombre = data.get('nombre', area.nombre).strip()
            area.piso = int(data.get('piso', area.piso))
            area.codigo_acceso = data.get('codigo_acceso', area.codigo_acceso).strip()
            
            dispositivo = data.get('dispositivo_biometrico') or data.get('dispositivo')
            if dispositivo and dispositivo in ['huella', 'Tarjeta', 'PIN']:
                area.dispositivo_biometrico = dispositivo
            
            nivel_seguridad = data.get('nivel_seguridad')
            if nivel_seguridad and nivel_seguridad in ['bajo', 'medio', 'alto']:
                area.nivel_seguridad = nivel_seguridad
            
            area.save()
            return JsonResponse({'success': True, 'message': 'Área actualizada correctamente'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    elif request.method == 'DELETE':
        try:
            area.delete()
            return JsonResponse({'success': True, 'message': 'Área eliminada correctamente'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)

    return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)


# ==================== EXPORTAR DATOS ====================

@login_required
def exportar_sedes(request, formato):
    """Exporta TODAS las ubicaciones (sedes + áreas)"""
    datos_exportar = []

    # Exportar sedes
    for sede in Sede.objects.all():
        datos_exportar.append({
            'ID': sede.id,
            'Tipo': 'Sede',
            'Código': sede.codigo_unico or f"SED-{sede.id:03d}",
            'Nombre': sede.nombre,
            'Sede Padre': '',
            'Piso': '',
            'Dirección': sede.direccion,
            'Dispositivo': sede.dispositivo_biometrico or '',
            'Seguridad': sede.nivel_seguridad or '',
            'Estado': 'Activo' if sede.activo else 'Inactivo',
        })

    # Exportar áreas
    for area in Area.objects.select_related('sede').all():
        datos_exportar.append({
            'ID': area.id,
            'Tipo': 'Área',
            'Código': area.codigo_acceso or f"ARE-{area.id:03d}",
            'Nombre': area.nombre,
            'Sede Padre': area.sede.nombre,
            'Piso': area.piso,
            'Dirección': area.sede.direccion,
            'Dispositivo': area.dispositivo_biometrico,
            'Seguridad': area.nivel_seguridad,
            'Estado': 'Activo',
        })

    if formato == 'json':
        response = HttpResponse(
            json.dumps(datos_exportar, ensure_ascii=False, indent=2),
            content_type='application/json'
        )
        response['Content-Disposition'] = 'attachment; filename="ubicaciones_export.json"'
        return response

    elif formato == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="ubicaciones_export.csv"'
        response.write('\ufeff')
        if datos_exportar:
            writer = csv.DictWriter(response, fieldnames=datos_exportar[0].keys())
            writer.writeheader()
            writer.writerows(datos_exportar)
        return response

    elif formato == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = "Ubicaciones"

        if datos_exportar:
            headers = list(datos_exportar[0].keys())
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = openpyxl.styles.Font(bold=True)

            for row_idx, row_data in enumerate(datos_exportar, 2):
                for col_idx, key in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ''))

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="ubicaciones_export.xlsx"'
        wb.save(response)
        return response

    return JsonResponse({'success': False, 'message': 'Formato no soportado'}, status=400)


# ==================== IMPORTAR DATOS ====================

@csrf_exempt
@login_required
def importar_sedes(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Método no permitido'}, status=405)

    archivo = request.FILES.get('archivo')
    if not archivo:
        return JsonResponse({'success': False, 'message': 'No se seleccionó ningún archivo'}, status=400)

    nombre_archivo = archivo.name.lower()
    datos_importados = []

    try:
        if nombre_archivo.endswith('.json'):
            contenido = archivo.read().decode('utf-8')
            datos_importados = json.loads(contenido)

        elif nombre_archivo.endswith('.csv'):
            contenido = archivo.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(contenido))
            for row in reader:
                datos_importados.append(dict(row))

        elif nombre_archivo.endswith(('.xlsx', '.xls')):
            wb = load_workbook(archivo)
            ws = wb.active
            headers = []
            for row in ws.iter_rows(values_only=True):
                if not headers:
                    headers = [str(cell) if cell else '' for cell in row]
                else:
                    row_data = {}
                    for idx, cell in enumerate(row):
                        if idx < len(headers) and headers[idx]:
                            row_data[headers[idx]] = cell
                    if any(row_data.values()):
                        datos_importados.append(row_data)
        else:
            return JsonResponse({'success': False, 'message': 'Formato no soportado. Use JSON, CSV o Excel'}, status=400)

        importados = 0
        errores = []

        for item in datos_importados:
            try:
                tipo = str(item.get('Tipo', item.get('tipo', 'Sede'))).strip()
                nombre = str(item.get('Nombre', item.get('nombre', ''))).strip()

                if not nombre:
                    errores.append(f"Fila sin nombre: {item}")
                    continue

                if tipo.lower() == 'sede':
                    sede, created = Sede.objects.get_or_create(
                        nombre=nombre,
                        defaults={
                            'direccion': item.get('Dirección', item.get('direccion', '')),
                            'codigo_unico': item.get('Código', item.get('codigo_unico', '')),
                            'dispositivo_biometrico': item.get('Dispositivo', item.get('dispositivo_biometrico', 'huella')),
                            'nivel_seguridad': item.get('Seguridad', item.get('nivel_seguridad', 'bajo')),
                            'activo': True
                        }
                    )
                    importados += 1

                elif tipo.lower() in ['área', 'area']:
                    sede_nombre = str(item.get('Sede Padre', item.get('sede_padre', ''))).strip()
                    try:
                        sede = Sede.objects.get(nombre=sede_nombre)
                    except Sede.DoesNotExist:
                        errores.append(f"Sede '{sede_nombre}' no encontrada para el área '{nombre}'")
                        continue

                    dispositivo = item.get('Dispositivo', item.get('dispositivo_biometrico', 'huella'))
                    if dispositivo not in ['huella', 'Tarjeta', 'PIN']:
                        dispositivo = 'huella'

                    nivel = str(item.get('Seguridad', item.get('nivel_seguridad', 'bajo'))).strip().lower()
                    if nivel not in ['bajo', 'medio', 'alto']:
                        nivel = 'bajo'

                    area, created = Area.objects.get_or_create(
                        sede=sede,
                        nombre=nombre,
                        defaults={
                            'piso': int(item.get('Piso', item.get('piso', 1)) or 1),
                            'codigo_acceso': item.get('Código', item.get('codigo_acceso', '')),
                            'dispositivo_biometrico': dispositivo,
                            'nivel_seguridad': nivel,
                        }
                    )
                    importados += 1

            except Exception as e:
                errores.append(f"Error importando '{item.get('Nombre', 'Desconocido')}': {str(e)}")

        return JsonResponse({
            'success': True,
            'message': f'Importación completada. Importados: {importados}, Errores: {len(errores)}',
            'errores': errores[:10],
            'importados': importados
        })

    except Exception as e:
        return JsonResponse({'success': False, 'message': f'Error al procesar archivo: {str(e)}'}, status=500)